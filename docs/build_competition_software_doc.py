from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_DIR = Path(__file__).resolve().parent
OUTPUT_FILE = OUTPUT_DIR / "NUEDC2026智能小车软件设计说明.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17365D"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
BLACK = "202124"
BORDER = "B7C4D3"

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, size=None, bold=None, color=None, italic=None,
                 ascii_font="Calibri", east_asia_font="Microsoft YaHei"):
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (
        ("top", top), ("start", start), ("bottom", bottom), ("end", end)
    ):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            cell.width = Inches(width / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color=MUTED)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("NUEDC2026 智能小车控制系统  |  软件设计说明")
    set_run_font(hr, size=9, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_field(fp)


def add_title_block(doc):
    doc.add_paragraph().paragraph_format.space_after = Pt(28)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.LEFT
    kicker.paragraph_format.space_after = Pt(4)
    r = kicker.add_run("比赛技术文档 · 软件部分")
    set_run_font(r, size=11, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run("NUEDC2026 智能小车控制系统")
    set_run_font(r, size=27, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(20)
    r = subtitle.add_run("任务调度、八路巡线、四轮闭环与终点控制软件设计说明")
    set_run_font(r, size=14, color=DARK_BLUE)

    metadata = [
        ("目标平台", "TI MSPM0G3507"),
        ("软件架构", "分层驱动 + 参数化任务模块 + 10 ms 周期状态机"),
        ("代码基线", "分支 2026H-LittleParsnip；提交 7015ff4；含当前工作区参数"),
        ("整理日期", "2026 年 7 月 30 日"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{label}：")
        set_run_font(r, size=10.5, bold=True, color=BLACK)
        r = p.add_run(value)
        set_run_font(r, size=10.5, color=BLACK)

    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    add_callout(
        doc,
        "文档定位",
        "本文以当前工程源码为唯一依据，重点说明可直接写入竞赛报告的软件架构、"
        "控制算法、状态机流程、关键参数和可靠性设计。文中“当前值”均对应生成文档时的磁盘版本。"
    )
    doc.add_page_break()


def add_callout(doc, label, text, fill=CALLOUT, trailing_space=True):
    table = doc.add_table(rows=1, cols=1)
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color="D8E0EA", size="6")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{label}  ")
    set_run_font(r, size=10.5, bold=True, color=DARK_BLUE)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=BLACK)
    if trailing_space:
        after = doc.add_paragraph()
        after.paragraph_format.space_after = Pt(2)
    else:
        # Word requires a paragraph after a table at the document end.
        # Keep it at 1 pt so it cannot create an otherwise blank final page.
        after = doc.add_paragraph()
        after.paragraph_format.space_before = Pt(0)
        after.paragraph_format.space_after = Pt(0)
        after.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        after.paragraph_format.line_spacing = Pt(1)
        run = after.add_run("")
        set_run_font(run, size=1)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r)


def add_numbers(doc, items):
    numbering = doc.part.numbering_part.element
    decimal_abstract_id = None
    for abstract in numbering.findall(qn("w:abstractNum")):
        for level in abstract.findall(qn("w:lvl")):
            num_fmt = level.find(qn("w:numFmt"))
            if (num_fmt is not None and
                    num_fmt.get(qn("w:val")) == "decimal"):
                decimal_abstract_id = abstract.get(qn("w:abstractNumId"))
                break
        if decimal_abstract_id is not None:
            break
    if decimal_abstract_id is None:
        raise RuntimeError("No decimal numbering definition found")

    existing_ids = [
        int(num.get(qn("w:numId")))
        for num in numbering.findall(qn("w:num"))
    ]
    num_id = max(existing_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_id = OxmlElement("w:abstractNumId")
    abstract_id.set(qn("w:val"), decimal_abstract_id)
    num.append(abstract_id)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)

    for item in items:
        p = doc.add_paragraph(style="List Number")
        p_pr = p._p.get_or_add_pPr()
        num_pr = p_pr.find(qn("w:numPr"))
        if num_pr is None:
            num_pr = OxmlElement("w:numPr")
            p_pr.append(num_pr)
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_node = OxmlElement("w:numId")
        num_id_node.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(num_id_node)
        r = p.add_run(item)
        set_run_font(r)


def add_code(doc, code):
    for line in code.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.18)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), LIGHT_GRAY)
        p_pr.append(shd)
        r = p.add_run(line if line else " ")
        set_run_font(
            r, size=9, color="263238",
            ascii_font="Consolas", east_asia_font="Microsoft YaHei"
        )
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def add_table(doc, headers, rows, widths_dxa, header_fill=LIGHT_BLUE,
              font_size=9.5, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(header))
        set_run_font(r, size=font_size, bold=True, color=NAVY)
    for row_values in rows:
        row = table.add_row()
        set_row_cant_split(row)
        cells = row.cells
        for index, value in enumerate(row_values):
            cell = cells[index]
            p = cell.paragraphs[0]
            if alignments:
                p.alignment = alignments[index]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    set_table_geometry(table, widths_dxa)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_static_contents(doc):
    add_heading(doc, "文档结构", 1)
    add_body(doc, "以下章节可直接作为比赛报告“软件设计”部分的主体框架。")
    items = [
        "软件设计目标与总体架构",
        "硬件抽象与底层驱动",
        "四轮速度闭环与车体运动控制",
        "八路灰度巡线算法",
        "任务选择、调度与执行",
        "一圈巡线任务状态机",
        "Task2 直线高速策略",
        "终点识别与平滑停车",
        "Task4 定距直行模块",
        "当前参数配置",
        "串口调试与测试方法",
        "软件特点、调参建议与版本说明",
    ]
    add_numbers(doc, items)
    add_callout(
        doc,
        "当前任务映射",
        "Task2、Task5、Task6 复用“一圈巡线”模块；Task4 使用定距直行模块；"
        "Task1、Task3 当前保留但未绑定执行模块。"
    )
    doc.add_page_break()


def build_document():
    doc = Document()
    style_document(doc)
    add_title_block(doc)
    add_static_contents(doc)

    add_heading(doc, "1. 软件设计目标与总体架构", 1)
    add_heading(doc, "1.1 设计目标", 2)
    add_body(
        doc,
        "软件面向四轮差速智能小车，要求在有限算力和强电磁干扰环境下完成稳定巡线、"
        "任务切换、定距行驶、终点停车和在线调试。设计重点不是把所有逻辑堆叠在 main.c，"
        "而是将硬件驱动、车体动作、巡线算法和比赛任务拆分为可独立调试的模块。"
    )
    add_bullets(doc, [
        "实时性：以 10 ms 为基本控制周期，电机闭环、巡线和任务状态机使用统一时间基准。",
        "可复用性：Task2、Task5、Task6 共用同一套一圈巡线状态机，仅替换参数表。",
        "可调性：速度、加减速时长、终点确认阈值、终点后前进距离均集中在 task_profiles.c。",
        "安全性：任务切换先停止旧模块；停车包含最小制动时间、停稳判定和超时故障分支。",
        "可观测性：UART 可输出灰度、电机目标/实测转速、编码器计数和 PWM，支持 VOFA+。"
    ])

    add_heading(doc, "1.2 分层架构", 2)
    add_table(
        doc,
        ["层级", "核心模块", "主要职责"],
        [
            ("比赛任务层", "task_manager / task_executor / task_profiles", "按键选题、模块分派、参数注入和运行互斥"),
            ("任务状态机层", "lap_task_control / straight_task_control", "一圈巡线、终点处理、定距直行与停车"),
            ("运动算法层", "line_tracking / angle_turn_control", "灰度误差计算、巡线 PID、定角转向"),
            ("车体控制层", "car_control", "将前进、转弯和差速动作换算为四轮目标 RPM"),
            ("执行器闭环层", "motor_control", "编码器测速、增量式 PID、PWM 与反向制动"),
            ("硬件接口层", "grayscale_sensor / SysConfig / DriverLib", "GPIO、定时器、UART、I²C 和传感器访问"),
        ],
        [1500, 3000, 4860],
    )
    add_body(
        doc,
        "数据流可以概括为：按键选择任务 → 执行器加载参数表 → 状态机给出线速度 → "
        "巡线模块计算左右差速 → 车体控制换算四轮目标转速 → 电机闭环输出 PWM。"
    )

    doc.add_page_break()
    add_heading(doc, "2. 硬件抽象与底层驱动", 1)
    add_heading(doc, "2.1 四轮布局与方向统一", 2)
    add_table(
        doc,
        ["电机", "车体位置", "前进符号", "说明"],
        [
            ("A", "右后轮", "-1", "右侧电机镜像安装，车体前进对应负目标 RPM"),
            ("B", "右前轮", "-1", "右侧电机镜像安装，车体前进对应负目标 RPM"),
            ("C", "左前轮", "+1", "左侧电机前进对应正目标 RPM"),
            ("D", "左后轮", "+1", "左侧电机前进对应正目标 RPM"),
        ],
        [1200, 1800, 1500, 4860],
        alignments=[
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
        ],
    )
    add_body(
        doc,
        "car_control.c 使用 forwardSign 将车体坐标系中的“向前”统一映射到各电机的物理方向，"
        "上层任务不需要关心 H 桥接线和电机镜像关系。"
    )

    add_heading(doc, "2.2 八路灰度硬件接口", 2)
    add_table(
        doc,
        ["信号", "当前引脚", "方向", "作用"],
        [
            ("AD0", "PB17", "输出", "CD4051 通道选择最低位"),
            ("AD1", "PA14", "输出", "CD4051 通道选择中间位"),
            ("AD2", "PA15", "输出", "CD4051 通道选择最高位"),
            ("OUT", "PB24", "输入（上拉+迟滞）", "读取所选灰度通道数字电平"),
        ],
        [1500, 1600, 2200, 4060],
    )
    add_body(
        doc,
        "每次选择通道后等待 50 μs，再以 5 μs 间隔读取 3 次并取多数值。"
        "这种做法兼顾 CD4051 建立时间和电机 PWM、编码器边沿带来的瞬态干扰。"
    )

    add_heading(doc, "3. 四轮速度闭环与车体运动控制", 1)
    add_heading(doc, "3.1 编码器测速", 2)
    add_body(
        doc,
        "四个电机均配置为编码器每转 13 脉冲、减速比 20、双倍解码，控制频率 100 Hz。"
        "电机输出轴每转对应计数为："
    )
    add_code(doc, "C_rev = encoderPpr × gearRatio × decodeMultiplier\n"
                  "      = 13 × 20 × 2 = 520 counts/rev")
    add_body(
        doc,
        "在 10 ms 采样周期内，目标 RPM 换算为目标计数："
        "C_target = RPM × C_rev / (60 × sampleRateHz)。"
    )

    add_heading(doc, "3.2 增量式 PID 与输出限幅", 2)
    add_body(
        doc,
        "motor_control.c 对每个轮子独立执行增量式 PID。当前四路参数一致："
        "Kp=5.0、Ki=1.0、Kd=0.2，正常驱动输出上限 99%，零速/超速反向制动上限 40%。"
    )
    add_code(doc, "Δu(k) = Kp[e(k)-e(k-1)] + Ki·e(k)\n"
                  "       + Kd[e(k)+e(k-2)-2e(k-1)]\n"
                  "u(k)  = sat(u(k-1)+Δu(k))")
    add_body(
        doc,
        "目标转速为零时，程序根据实测运动方向施加反向制动力，直至编码器计数进入死区。"
        "目标非零但实测速度显著超调时，也允许有限反向 PWM；这解决了弯道内轮设为低速后"
        "仍被车体拖动、无法形成有效差速的问题。"
    )
    add_callout(
        doc,
        "反向制动的意义",
        "传统单象限调速在超速时只能把 PWM 降为 0，轮子仍会滑行。当前实现允许最多 40% 的"
        "反向制动力，使低速内轮真正降速，从而提升急弯转向能力。"
    )

    add_heading(doc, "3.3 车体动作映射", 2)
    add_table(
        doc,
        ["动作", "左侧轮", "右侧轮", "用途"],
        [
            ("前进", "v", "v", "直线行驶"),
            ("前进左转", "v·p", "v", "左侧为内轮，p 为内轮百分比"),
            ("前进右转", "v", "v·p", "右侧为内轮"),
            ("原地左转", "-v", "+v", "定角转向"),
            ("原地右转", "+v", "-v", "定角转向"),
        ],
        [1800, 1800, 1800, 3960],
    )

    add_heading(doc, "4. 八路灰度巡线算法", 1)
    add_heading(doc, "4.1 传感器消抖与加权误差", 2)
    add_body(
        doc,
        "黑线有效电平配置为 1。八路传感器 X1～X8 从左到右赋予位置权重："
    )
    add_code(doc, "weights = {-30, -20, -15, 0, 0, 15, 20, 30}")
    add_body(
        doc,
        "多路同时有效时采用加权平均：error = Σ(weight_i × active_i) / Σ(active_i)。"
        "仅中央 X4/X5 有效时直接返回 0，避免中央边缘抖动造成左右摆动；误差绝对值不超过 5"
        "时也按居中处理。每一路数字状态需连续 2 个采样周期一致后才更新。"
    )

    add_heading(doc, "4.2 丢线恢复与巡线 PID", 2)
    add_body(
        doc,
        "若所有通道均未检测到黑线，程序依据上一有效误差方向输出 ±30 的搜索误差，"
        "使车辆沿原偏转方向继续寻找轨迹。巡线控制器当前参数如下："
    )
    add_table(
        doc,
        ["参数", "当前值", "作用"],
        [
            ("Kp", "3.0", "决定偏差对应的主要差速强度"),
            ("Ki", "0.01", "补偿长期偏差；居中或丢线时积分清零"),
            ("Kd", "0.0", "当前关闭，避免数字误差跳变放大"),
            ("误差死区", "±5", "减少中心附近左右抖动"),
            ("差速死区", "8 RPM", "修正量较小时保持直行"),
            ("误差滤波 α", "1.00", "等效不滤波；保留接口便于后续实验"),
        ],
        [2000, 1800, 5560],
    )
    add_body(
        doc,
        "PID 输出被限制在 ±baseSpeedRpm。修正量为负时执行前进左转，为正时执行前进右转；"
        "内轮目标速度为 baseSpeedRpm - |PID offset|，最低保留为 1%。"
    )
    add_callout(
        doc,
        "调参结论",
        "数字灰度误差具有阶跃特性。过强低通会增加相位滞后并导致车体越过中心后仍继续转向，"
        "因此当前 α=1.00；平滑性主要依靠传感器消抖、误差死区、适度 Kp 和速度斜坡。"
    )

    add_heading(doc, "5. 任务选择、调度与执行", 1)
    add_heading(doc, "5.1 按键管理", 2)
    add_body(
        doc,
        "TASK、TASK_START 和 STATUS 三个按键由 task_manager.c 管理。按键电平连续稳定 3 个"
        "10 ms 周期后才确认状态变化。TASK 键按 Task1→Task6 循环选择，TASK_START 产生一次性"
        "启动事件，OLED 用于显示当前任务号和运行计时。"
    )
    add_heading(doc, "5.2 执行器映射与互斥", 2)
    add_table(
        doc,
        ["任务", "执行模块", "当前状态"],
        [
            ("Task1", "无", "保留任务号，当前不驱动车辆"),
            ("Task2", "LapTaskControl", "一圈巡线；直线高速 280 RPM"),
            ("Task3", "无", "保留任务号，当前不驱动车辆"),
            ("Task4", "StraightTaskControl", "编码器定距直行 2000 mm"),
            ("Task5", "LapTaskControl", "一圈巡线，参数独立"),
            ("Task6", "LapTaskControl", "一圈巡线，参数独立"),
        ],
        [1400, 2800, 5160],
    )
    add_body(
        doc,
        "任务切换时执行器先复位一圈模块和直行模块，再启动新模块。任一时刻只有一个状态机拥有"
        "车辆控制权，从结构上避免两个任务同时写入电机目标速度。"
    )

    add_heading(doc, "6. 一圈巡线任务状态机", 1)
    add_body(
        doc,
        "一圈任务使用 FOLLOW → FINISH_ADVANCE → DECELERATE → BRAKE → DONE/FAULT 的状态流。"
        "每个状态只处理一类职责，便于分别调试终点识别、距离推进和停车。"
    )
    add_table(
        doc,
        ["状态", "主要输出", "转移条件"],
        [
            ("FOLLOW", "启动加速、巡线 PID、直线提速、持续检测十字路口", "终点确认后进入 FINISH_ADVANCE"),
            ("FINISH_ADVANCE", "关闭巡线；锁定终点前转弯比例；按当前速度继续前进", "累计编码器达到 finishAdvanceMm"),
            ("DECELERATE", "从终点入口速度线性降至 decelerationEndRpm", "减速周期完成"),
            ("BRAKE", "四轮目标归零并主动制动", "最小制动时间后四轮均进入死区"),
            ("DONE", "保持任务完成状态", "等待下一次任务操作"),
            ("FAULT", "紧急滑行停止", "制动等待超过 brakeTimeoutSamples"),
        ],
        [1700, 4660, 3000],
        font_size=9.2,
    )

    add_heading(doc, "6.1 十字路口识别", 2)
    add_body(
        doc,
        "若八路全部有效，程序立即认为检测到十字路口；否则当有效通道数达到参数"
        " intersectionActiveThreshold 并连续满足 intersectionConfirmSamples 次后确认。"
        "该策略兼顾完整横线的快速响应和部分遮挡情况下的抗误判能力。"
    )

    add_heading(doc, "6.2 终点抖动抑制", 2)
    add_body(
        doc,
        "终点横线会造成多个通道同时跳变。确认终点后，程序在同一周期立即关闭巡线 PID，"
        "调用 LineTracking_reset 清空积分、上次误差、误差滤波和传感器消抖历史，且不再执行"
        "该周期的 LineTracking_update。这样横线不会继续被解释为左右转向误差。"
    )
    add_body(
        doc,
        "关闭巡线前，状态机会保存当时的速度指令、转弯方向和内轮百分比。终点后不突然改为"
        "等速直行，而是保持最后的左右轮速度比例。例如终点前为外轮 100、内轮 40 RPM，"
        "减速时按 80/32、40/16、10/4 变化，保持路径曲率连续。"
    )

    add_heading(doc, "7. Task2 直线高速策略", 1)
    add_body(
        doc,
        "Task2 将 cruiseRpm 作为弯道/普通巡线速度，将 straightRpm 作为直行目标速度。"
        "基础启动斜坡完成后，如果上一周期巡线动作是 CAR_CONTROL_FORWARD，则逐步增加"
        "直线提速计数；出现左转或右转修正时逐步减少该计数。"
    )
    add_code(doc, "v = cruiseRpm + (straightRpm - cruiseRpm) × n / N\n"
                  "n: 当前直线提速计数；N: straightAccelerationSamples")
    add_body(
        doc,
        "当前 Task2 配置为弯道 150 RPM、直线 280 RPM，直线提速过渡 50×10 ms=0.5 s。"
        "这种“按轨迹状态调度速度”的方法在直线段提高平均速度，同时避免以 280 RPM 强行进入弯道。"
    )

    add_heading(doc, "8. 终点速度继承与平滑停车", 1)
    add_body(
        doc,
        "终点入口速度不是固定采用配置巡航速度，而是读取 LineTracking 当前实际下发的"
        " baseSpeedRpm。这样车辆若在加速未完成或刚从直线高速回落时到达终点，不会先跳到"
        "配置速度再减速。"
    )
    add_code(doc, "v_finish_start = LineTracking_getSpeed()\n"
                  "v(k) = v_finish_start\n"
                  "       - (v_finish_start - v_end) × k / N_decel")
    add_body(
        doc,
        "同时限定 v_end ≤ v_finish_start，避免低速经过终点时“减速阶段”反而重新加速。"
        "终点后前进距离由四轮编码器绝对计数总和换算，达到 finishAdvanceMm 后才进入减速。"
    )

    add_heading(doc, "9. Task4 定距直行模块", 1)
    add_body(
        doc,
        "Task4 不依赖灰度传感器，使用四轮编码器累计计数控制距离。距离换算采用轮径和"
        "四个电机每圈计数总和："
    )
    add_code(doc, "C_distance = distance_mm × ΣC_rev / (π × wheelDiameter_mm)")
    add_body(
        doc,
        "行驶过程分为启动线性加速、中段匀速和按剩余距离连续减速三部分。进入最后"
        " decelerationDistanceMm 后，目标速度随剩余距离线性下降；达到目标距离后进入"
        "主动制动与停稳确认。"
    )

    add_heading(doc, "10. 当前任务参数配置", 1)
    add_heading(doc, "10.1 一圈巡线任务", 2)
    add_table(
        doc,
        ["参数", "Task2", "Task5", "Task6", "单位/说明"],
        [
            ("cruiseRpm", "150", "130", "130", "RPM，弯道/普通巡线速度"),
            ("straightRpm", "280", "130", "130", "RPM，直行目标速度"),
            ("accelerationSamples", "80", "600", "600", "10 ms/计数"),
            ("straightAccelerationSamples", "50", "0", "0", "直线速度过渡周期"),
            ("intersectionActiveThreshold", "4", "4", "4", "有效通道数"),
            ("intersectionConfirmSamples", "2", "1", "1", "连续确认次数"),
            ("finishAdvanceMm", "100", "5", "5", "终点后继续前进距离"),
            ("wheelDiameterMm", "65", "65", "65", "有效轮径"),
            ("decelerationEndRpm", "10", "5", "5", "减速结束速度"),
            ("decelerationSamples", "30", "200", "200", "减速周期"),
            ("brakeMinimumSamples", "15", "15", "15", "最小制动周期"),
            ("brakeTimeoutSamples", "100", "100", "100", "制动超时周期"),
        ],
        [2900, 1100, 1100, 1100, 3160],
        font_size=8.8,
        alignments=[
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
        ],
    )

    add_heading(doc, "10.2 Task4 定距直行", 2)
    add_table(
        doc,
        ["参数", "当前值", "说明"],
        [
            ("targetDistanceMm", "2000 mm", "目标行驶距离"),
            ("wheelDiameterMm", "65 mm", "距离换算有效轮径"),
            ("cruiseRpm", "150 RPM", "中段匀速目标"),
            ("accelerationSamples", "300", "启动加速 3.0 s"),
            ("decelerationDistanceMm", "300 mm", "终点前减速区长度"),
            ("decelerationEndRpm", "5 RPM", "到达终点前的最低目标速度"),
            ("brakeMinimumSamples", "15", "至少制动 0.15 s"),
            ("brakeTimeoutSamples", "100", "1.0 s 未停稳则故障"),
        ],
        [3000, 1900, 4460],
    )

    add_heading(doc, "11. 串口调试与测试方法", 1)
    add_heading(doc, "11.1 常用命令", 2)
    add_table(
        doc,
        ["命令", "功能", "典型用途"],
        [
            ("F/B/L/R rpm", "前进、后退、原地左转、原地右转", "检查车轮方向映射"),
            ("FL/FR rpm inner%", "前进差速左转/右转", "验证内外轮速度比例"),
            ("WA/WB/WC/WD rpm", "单独控制某个电机", "定位单轮硬件或 PID 问题"),
            ("I speed / I1 / I0", "启动巡线、带调试巡线、停止巡线", "观察误差和八路状态"),
            ("G / G1 / G0", "灰度单次/连续/关闭输出", "校验 X1～X8 极性和顺序"),
            ("M / M1 / M0", "电机单次/连续/关闭输出", "观察目标、实测、计数和 PWM"),
            ("U3 text", "向 UART3 发送文本并自动追加 CRLF", "联调外部串口设备"),
        ],
        [1900, 3300, 4160],
        font_size=9.1,
    )
    add_heading(doc, "11.2 电机数据帧", 2)
    add_body(
        doc,
        "VOFA+ FireWater 帧以 motor: 开头，每个电机依次输出 targetRpm、speedRpm、"
        "100 ms 编码器累计计数和 pwmPercent。A、B 前进时目标为负，C、D 前进时目标为正，"
        "这是安装方向符号导致的正常现象。"
    )
    add_code(doc, "motor:A目标,A实测,A计数,A PWM,\n"
                  "      B目标,B实测,B计数,B PWM,\n"
                  "      C目标,C实测,C计数,C PWM,\n"
                  "      D目标,D实测,D计数,D PWM")
    add_numbers(doc, [
        "先使用 WA/WB/WC/WD 验证每个轮子的方向、编码器符号和闭环稳定性。",
        "使用 G1 确认黑线对应 1，且 X1 位于车体左侧、X8 位于右侧。",
        "使用 I1 低速观察 line:error,pid,X1…X8，确认左偏误差为负、右偏为正。",
        "最后运行 Task2/5/6，使用 M1 检查弯道内轮是否出现合理反向制动。"
    ])

    add_heading(doc, "12. 软件特点、调参建议与版本说明", 1)
    add_heading(doc, "12.1 可写入比赛报告的软件特点", 2)
    add_bullets(doc, [
        "参数化任务复用：一套状态机通过不同 Profile 支持多个比赛任务，减少重复代码。",
        "双层闭环：巡线 PID 决定车体差速，四个电机 PID 分别保证轮速跟踪。",
        "非零目标反向制动：解决低速内轮被拖动导致差速失效的问题。",
        "直线自适应提速：根据巡线动作自动在弯道速度与直线高速之间平滑切换。",
        "终点控制连续性：终点时冻结最后转向比例、继承当前速度并线性减速。",
        "可观测调试接口：灰度、电机、角度和 UART3 数据均可通过串口在线观测。"
    ])

    add_heading(doc, "12.2 推荐调参顺序", 2)
    add_numbers(doc, [
        "先标定四个电机方向符号、编码器符号和单轮速度 PID。",
        "低速校验八路灰度极性、物理顺序和中央死区。",
        "保持直线速度较低，仅调巡线 Kp；确认能过最急弯后再微调 Ki。",
        "验证反向制动上限，确保内轮能降速但不会产生明显冲击。",
        "分别调整 cruiseRpm、straightRpm 和 straightAccelerationSamples。",
        "最后标定轮径、终点后前进距离、减速周期与停稳判定。"
    ])

    add_heading(doc, "12.3 当前版本注意事项", 2)
    add_bullets(doc, [
        "当前 LINE_ERROR_FILTER_ALPHA=1.00，误差低通接口存在但等效关闭；比赛文档不应描述为已启用 30% 低通。",
        "Task1、Task3 目前未绑定控制模块；如赛题需要，应新增对应模块或明确作为保留选项。",
        "Task4 当前参数为 2000 mm，但部分历史日志/注释仍写 1500 mm，提交比赛材料前应统一文字。",
        "task_executor.c 中 Stopwatch_stop() 当前被注释，若需要任务完成后冻结计时，应恢复并验证调用时机。",
        "参数经常在 CCS 中现场调整，提交前应重新核对 task_profiles.c 与本表是否一致。"
    ])

    add_heading(doc, "附录 A：关键源文件职责", 1)
    add_table(
        doc,
        ["文件", "职责"],
        [
            ("main.c", "硬件对象配置、系统初始化、10 ms 调度入口、UART 调试命令"),
            ("motor_control.c/.h", "单电机编码器测速、增量 PID、PWM、滑行与反向制动"),
            ("car_control.c/.h", "车体动作到四轮目标 RPM 的映射及最后运动状态保存"),
            ("grayscale_sensor.c/.h", "CD4051 八路通道选择、延时、三次多数采样"),
            ("line_tracking.c/.h", "传感器消抖、加权误差、丢线恢复和巡线 PID"),
            ("task_manager.c/.h", "按键消抖、任务选择、启动和状态事件"),
            ("task_executor.c/.h", "任务到模块的映射、互斥启动和周期更新"),
            ("task_profiles.c/.h", "Task2/4/5/6 当前比赛参数"),
            ("lap_task_control.c/.h", "一圈巡线、终点识别、速度继承、比例锁定和停车状态机"),
            ("straight_task_control.c/.h", "编码器定距直行、距离减速和停稳状态机"),
            ("stopwatch.c/.h", "任务计时及 OLED 显示时间来源"),
        ],
        [3000, 6360],
    )

    add_callout(
        doc,
        "结论",
        "当前软件已形成“硬件驱动—单轮闭环—车体差速—巡线算法—任务状态机—参数配置”的"
        "完整链路。比赛报告可重点突出模块复用、双层闭环、直线提速、反向制动和终点连续控制。",
        trailing_space=False,
    )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_FILE)
    print(OUTPUT_FILE)


if __name__ == "__main__":
    build_document()
